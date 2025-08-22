from flask import Flask, request, jsonify
from flask_cors import CORS
import pandas as pd
import os
from werkzeug.utils import secure_filename
import tempfile
from datetime import datetime, timedelta
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import numbers

app = Flask(__name__)
CORS(app)

UPLOAD_FOLDER = tempfile.gettempdir()
ALLOWED_EXTENSIONS = {'xlsx', 'xls'}

app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024

holidays = pd.to_datetime([
    "2022-01-01", "2022-01-02","2022-01-03", "2022-01-04","2022-01-20",
    "2022-03-08","2022-03-20", "2022-03-21", "2022-03-22", "2022-03-23", "2022-03-24","2022-03-25",  
    "2022-05-02", "2022-05-03",  "2022-05-09", "2022-05-28", "2022-05-30", "2022-06-15","2022-06-26","2022-06-27","2022-07-09", "2022-07-10","2022-07-11","2022-07-12", "2022-11-08","2022-11-09",
    "2022-12-31", 
    "2023-01-01", "2023-01-02","2023-01-03","2023-01-04",         
    "2023-01-20", "2023-03-08",                     
    "2023-03-20", "2023-03-21", "2023-03-22", "2023-03-23", "2023-03-24", 
    "2023-04-21", "2023-04-22", "2023-04-24",      
    "2023-05-09",                       
    "2023-05-28", "2023-05-29",                     
    "2023-06-15",                       
    "2023-06-26",                       
    "2023-06-28", "2023-06-29",         
    "2023-11-08",                       
    "2023-11-09",                                            
    "2023-12-31",
    "2024-01-01", "2024-01-02", "2024-01-03", "2024-01-04", "2024-01-05",  
    "2024-01-20",                                                        
    "2024-02-07",                                                         
    "2024-03-08",                                                         
    "2024-03-20", "2024-03-21", "2024-03-22", "2024-03-23", "2024-03-24", 
    "2024-03-25", "2024-03-26",                                           
    "2024-04-10", "2024-04-11",                                           
    "2024-05-09",                                                         
    "2024-05-28",                                                         
    "2024-06-15",                                                         
    "2024-06-16", "2024-06-17",                                          
    "2024-06-18", "2024-06-19",                                           
    "2024-06-26",                                                         
    "2024-11-08",                                                        
    "2024-11-09", "2024-11-11",                                          
    "2024-12-23", "2024-12-31",
    "2025-01-01", "2025-01-02", "2025-01-03",           
    "2025-01-20", "2025-01-29",                         
    "2025-03-08", "2025-03-10",                          
    "2025-03-20", "2025-03-21", "2025-03-22", "2025-03-23", "2025-03-24",  
    "2025-03-25", "2025-03-26", "2025-03-30", "2025-03-31",  
    "2025-04-01",                          
    "2025-05-09",                                      
    "2025-05-28",                                       
    "2025-06-06", "2025-06-07", "2025-06-09",           
    "2025-06-15",                                       
    "2025-06-16",                          
    "2025-06-26",                                       
    "2025-11-08",                                      
    "2025-11-09", "2025-11-10", "2025-11-11",           
    "2025-12-31"
])

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def is_weekend_or_holiday(date):
    if date.weekday() >= 5:
        return True
    if date in holidays:
        return True
    return False

def adjust_column_widths(worksheet):
    for col in worksheet.columns:
        max_length = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                cell_length = len(str(cell.value))
                if cell_length > max_length:
                    max_length = cell_length
        adjusted_width = max_length + 2
        worksheet.column_dimensions[col_letter].width = adjusted_width

class CMSChecker:
    def __init__(self, excel_file_path):
        self.excel_file = excel_file_path
        self.df = pd.read_excel(excel_file_path)
    
    def access_termination(self):
        try:
            results = []
            condition_nonopen = (
                self.df['Termination Date'].notna() & 
                self.df['Group Name'].notna() & 
                (self.df['Group Name'].astype(str).str.strip() != '') &
                (self.df['Deleted On Db'].astype(str).str.lower() != 'yes') &
                (self.df['User Status'].astype(str).str.upper() != 'OPEN') &
                (self.df['Hr Code'].astype(str).str.upper().str.strip() != 'S')
            )
            filtered_nonopen_df = self.df[condition_nonopen]
            df_copy = self.df.copy()
            df_copy['Termination Date'] = pd.to_datetime(df_copy['Termination Date'], errors='coerce')
            filtered_open_df = df_copy[
                df_copy['Termination Date'].notna() &
                (df_copy['User Status'].str.upper() == 'OPEN')
            ]
            df_compare = self.df.copy()
            df_compare['Termination Date'] = pd.to_datetime(df_compare['Termination Date'], dayfirst=True, errors='coerce').dt.normalize()
            df_compare['User Lock Date'] = pd.to_datetime(df_compare['User Lock Date'], dayfirst=True, errors='coerce')
            def has_lock_problem(row):
                term_date = row['Termination Date']
                lock_date = row['User Lock Date']
                if pd.isna(term_date) or pd.isna(lock_date):
                    return False
                if term_date.date() == lock_date.date():
                    return False
                day_diff = (lock_date.date() - term_date.date()).days
                if day_diff <= 1:
                    return False
                else:
                    if is_weekend_or_holiday(term_date):
                        next_workday = term_date + timedelta(days=1)
                        while is_weekend_or_holiday(next_workday):
                            next_workday += timedelta(days=1)
                        if lock_date.date() > next_workday.date():
                            return True
                        else:
                            return False
                    else:
                        return True
            problem_df = df_compare[df_compare.apply(has_lock_problem, axis=1)]
            results.append({
                'status': 'success',
                'title': 'Non-Open Users',
                'message': f'{len(filtered_nonopen_df)} istifadəçi tapıldı'
            })
            if len(filtered_open_df) > 0:
                results.append({
                    'status': 'warning',
                    'title': 'Open Users (İxtisar edilmiş)',
                    'message': f'{len(filtered_open_df)} istifadəçi hələ OPEN statusundadır'
                })
            if len(problem_df) > 0:
                results.append({
                    'status': 'error',
                    'title': 'Lock Problem Users',
                    'message': f'{len(problem_df)} istifadəçinin lock tarixində problem var'
                })
            return {'success': True, 'details': results, 'data': {
                'non_open': filtered_nonopen_df.to_dict('records'),
                'open_users': filtered_open_df.to_dict('records'),
                'lock_problems': problem_df.to_dict('records')
            }}
        except Exception as e:
            return {'success': False, 'error': str(e)}
    
    def expiry_date_monitoring(self):
        try:
            filtered_df = self.df[
                (self.df['User Status'] == 'OPEN') &
                (self.df['User Expiry Date'].isna()) &
                (self.df['Hr Code'].str.startswith(('T', 'V'), na=False))
            ]
            results = [{
                'status': 'warning' if len(filtered_df) > 0 else 'success',
                'title': 'Expiry Date Check',
                'message': f'{len(filtered_df)} istifadəçinin expiry date-i yoxdur'
            }]
            return {'success': True, 'details': results, 'data': filtered_df.to_dict('records')}
        except Exception as e:
            return {'success': False, 'error': str(e)}
    
    def user_workplace_match(self):
        try:
            filtered_df = self.df[
                self.df['User Name'].str.startswith(('BO', 'HO'), na=False) &
                (self.df['User Status'] == 'OPEN') &
                (self.df['Current Position'].astype(str).str.strip() == '/  /  /') &
                self.df['Hr Code'].str.startswith('T', na=False)
            ]
            filtered_df = filtered_df.drop_duplicates(subset=['Hr Code'], keep='first')
            results = [{
                'status': 'warning' if len(filtered_df) > 0 else 'success',
                'title': 'User-Workplace Match',
                'message': f'{len(filtered_df)} uyğunsuzluq tapıldı'
            }]
            return {'success': True, 'details': results, 'data': filtered_df.to_dict('records')}
        except Exception as e:
            return {'success': False, 'error': str(e)}
    
    def vacation_monitoring(self):
        try:
            df_copy = self.df.copy()
            df_copy['Vacation End Date'] = pd.to_datetime(df_copy['Vacation End Date'], format='%d/%m/%Y', errors='coerce')
            current_date = datetime.now()
            df_with_vacation = df_copy[df_copy['Vacation End Date'].notna()].copy()
            df_with_vacation['Days Until Vacation End'] = (df_with_vacation['Vacation End Date'] - current_date).dt.days
            df_over_60_days = df_with_vacation[df_with_vacation['Days Until Vacation End'] > 60]
            df_result = df_over_60_days[df_over_60_days['User Status'].str.upper() == 'OPEN']
            results = [{
                'status': 'success',
                'title': 'Vacation Monitoring',
                'message': f'{len(df_result)} istifadəçinin məzuniyyəti 60 gündən çox qalıb'
            }]
            return {'success': True, 'details': results, 'data': df_result.to_dict('records')}
        except Exception as e:
            return {'success': False, 'error': str(e)}

class FIMIChecker:
    def __init__(self, excel_file_path):
        self.excel_file = excel_file_path
        self.df = pd.read_excel(excel_file_path)
    
    def access_termination(self):
        try:
            results = []
            condition_nonactive = (
                self.df['TERMINATION_DATE'].notna() & 
                self.df['ACCESS_GROUP'].notna() & 
                (self.df['ACCESS_GROUP'].astype(str).str.strip() != '') & 
                (~self.df['ACCESS_GROUP'].astype(str).str.lower().str.contains('blocked', na=False)) &
                (self.df['USER_STATUS'].astype(str).str.upper() != 'ACTIVE') &
                (self.df['AD_STATUS'].astype(str).str.upper() != 'TRUE') &
                (self.df['HR_CODE'].astype(str).str.upper().str.strip() != 'S')
            )
            filtered_nonactive_df = self.df[condition_nonactive]
            df_copy = self.df.copy()
            df_copy['TERMINATION_DATE'] = pd.to_datetime(df_copy['TERMINATION_DATE'], errors='coerce')
            filtered_active_df = df_copy[
                df_copy['TERMINATION_DATE'].notna() &
                ((df_copy['USER_STATUS'].astype(str).str.upper() == 'ACTIVE') |
                 (df_copy['AD_STATUS'].astype(str).str.upper() == 'TRUE'))
            ]
            results.append({
                'status': 'success',
                'title': 'Non-Active Users',
                'message': f'{len(filtered_nonactive_df)} istifadəçi tapıldı'
            })
            if len(filtered_active_df) > 0:
                results.append({
                    'status': 'warning',
                    'title': 'Active Users (İxtisar edilmiş)',
                    'message': f'{len(filtered_active_df)} istifadəçi hələ ACTIVE statusundadır'
                })
            return {'success': True, 'details': results, 'data': {
                'non_active': filtered_nonactive_df.to_dict('records'),
                'active_users': filtered_active_df.to_dict('records')
            }}
        except Exception as e:
            return {'success': False, 'error': str(e)}
    
    def user_workplace_match(self):
        try:
            filtered_df = self.df[
                self.df['USERNAME'].str.startswith(('BO', 'HO'), na=False) &
                (self.df['USER_STATUS'] == 'Active') &
                (self.df['CURRENT_POSITION'].isna() | 
                 (self.df['CURRENT_POSITION'].astype(str).str.strip() == '') |
                 (self.df['CURRENT_POSITION'].astype(str).str.strip() == '/  /  /')) &
                self.df['HR_CODE'].str.startswith('T', na=False)
            ]
            filtered_df = filtered_df.drop_duplicates(subset=['HR_CODE'], keep='first')
            results = [{
                'status': 'warning' if len(filtered_df) > 0 else 'success',
                'title': 'User-Workplace Match',
                'message': f'{len(filtered_df)} uyğunsuzluq tapıldı'
            }]
            return {'success': True, 'details': results, 'data': filtered_df.to_dict('records')}
        except Exception as e:
            return {'success': False, 'error': str(e)}
    
    def vacation_monitoring(self):
        try:
            df_copy = self.df.copy()
            df_copy['VACATION_END_DATE'] = pd.to_datetime(df_copy['VACATION_END_DATE'], format='%d/%m/%Y', errors='coerce')
            current_date = datetime.now()
            df_with_vacation = df_copy[df_copy['VACATION_END_DATE'].notna()].copy()
            df_with_vacation['Days_Until_Vacation_End'] = (df_with_vacation['VACATION_END_DATE'] - current_date).dt.days
            df_over_60_days = df_with_vacation[df_with_vacation['Days_Until_Vacation_End'] > 60]
            df_result = df_over_60_days[df_over_60_days['USER_STATUS'].str.upper() == 'ACTIVE']
            results = [{
                'status': 'success',
                'title': 'Vacation Monitoring',
                'message': f'{len(df_result)} istifadəçinin məzuniyyəti 60 gündən çox qalıb'
            }]
            return {'success': True, 'details': results, 'data': df_result.to_dict('records')}
        except Exception as e:
            return {'success': False, 'error': str(e)}

class KapitalCardChecker:
    def __init__(self, excel_file_path):
        self.excel_file = excel_file_path
        self.df = pd.read_excel(excel_file_path)
    
    def access_termination(self):
        try:
            df_copy = self.df.copy()
            df_copy['TERMINATION_DATE'] = pd.to_datetime(df_copy['TERMINATION_DATE'], dayfirst=True, errors='coerce')
            df_copy['LAST_LOGON_DATE'] = pd.to_datetime(df_copy['LAST_LOGON_DATE'], errors='coerce')
            filtered_df_1 = df_copy[(df_copy['TERMINATION_DATE'].notna()) & (df_copy['ENABLED'] == True)]
            df_copy['TERMINATION_DATE_DATEONLY'] = df_copy['TERMINATION_DATE'].dt.date
            df_copy['LAST_LOGON_DATE_DATEONLY'] = df_copy['LAST_LOGON_DATE'].dt.date
            filtered_df_2 = df_copy[(df_copy['ENABLED'] == False) & 
                                   (df_copy['LAST_LOGON_DATE_DATEONLY'] > df_copy['TERMINATION_DATE_DATEONLY'])]
            results = []
            if len(filtered_df_1) > 0:
                results.append({
                    'status': 'warning',
                    'title': 'Terminated but Enabled',
                    'message': f'{len(filtered_df_1)} istifadəçi ixtisar edilib amma hələ aktivdir'
                })
            if len(filtered_df_2) > 0:
                results.append({
                    'status': 'error',
                    'title': 'Logon After Termination',
                    'message': f'{len(filtered_df_2)} istifadəçi ixtisardan sonra daxil olub'
                })
            return {'success': True, 'details': results, 'data': {
                'terminated_enabled': filtered_df_1.to_dict('records'),
                'logon_after_term': filtered_df_2.to_dict('records')
            }}
        except Exception as e:
            return {'success': False, 'error': str(e)}
    
    def vendor_access_monitoring(self):
        try:
            filtered_df = self.df[
                (self.df['ENABLED'] == True) &
                (self.df['PO_BOX'] == 'V') &
                (self.df['DESCRIPTION'].astype(str).str.contains("SD-", na=False)) &
                (self.df['ACCOUNT_EXPIRATION_DATE'].isna())
            ]
            results = [{
                'status': 'warning' if len(filtered_df) > 0 else 'success',
                'title': 'Vendor Access Check',
                'message': f'{len(filtered_df)} vendor hesabının müddəti təyin edilməyib'
            }]
            return {'success': True, 'details': results, 'data': filtered_df.to_dict('records')}
        except Exception as e:
            return {'success': False, 'error': str(e)}
    
    def procedure_alignment_check(self):
        try:
            results = []
            excluded_values = ['S', 'X', 'MR', 'RM']
            df_main = self.df[~self.df['PO_BOX'].astype(str).str.strip().str.upper().isin(excluded_values)]
            not_required_df = df_main[
                (df_main['PASSWORD_NOT_REQUIRED'].notna()) &
                (df_main['PASSWORD_NOT_REQUIRED'].astype(str).str.strip() != '') &
                (df_main['PASSWORD_NOT_REQUIRED'] == True)
            ]
            df_main['PASSWORD_LAST_SET'] = pd.to_datetime(df_main['PASSWORD_LAST_SET'], errors='coerce')
            today = pd.to_datetime(datetime.today().date())
            df_main['DAY_DIFF'] = (today - df_main['PASSWORD_LAST_SET']).dt.days
            password_check_df = df_main[
                (df_main['DAY_DIFF'] >= 60) &
                (df_main['ENABLED'] == True) &
                (df_main['PASSWORD_EXPIRED'] == False) &
                (df_main['PASSWORD_NEVER_EXPIRES'] == False)
            ]
            if len(not_required_df) > 0:
                results.append({
                    'status': 'error',
                    'title': 'Password Not Required',
                    'message': f'{len(not_required_df)} hesabda şifrə tələb edilmir'
                })
            if len(password_check_df) > 0:
                results.append({
                    'status': 'warning',
                    'title': 'Old Passwords',
                    'message': f'{len(password_check_df)} hesabın şifrəsi 60 gündən köhnədir'
                })
            return {'success': True, 'details': results, 'data': {
                'password_not_required': not_required_df.to_dict('records'),
                'old_passwords': password_check_df.to_dict('records')
            }}
        except Exception as e:
            return {'success': False, 'error': str(e)}

class KapitalhoChecker:
    def __init__(self, excel_file_path):
        self.excel_file = excel_file_path
        self.df = pd.read_excel(excel_file_path)
    
    def access_termination(self):
        try:
            df_copy = self.df.copy()
            df_copy['TERMINATION_DATE'] = pd.to_datetime(df_copy['TERMINATION_DATE'], dayfirst=True, errors='coerce')
            df_copy['LAST_LOGON_DATE'] = pd.to_datetime(df_copy['LAST_LOGON_DATE'], errors='coerce')
            filtered_df_1 = df_copy[(df_copy['TERMINATION_DATE'].notna()) & (df_copy['ENABLED'] == True)]
            df_copy['TERMINATION_DATE_DATEONLY'] = df_copy['TERMINATION_DATE'].dt.date
            df_copy['LAST_LOGON_DATE_DATEONLY'] = df_copy['LAST_LOGON_DATE'].dt.date
            filtered_df_2 = df_copy[(df_copy['ENABLED'] == False) & 
                                   (df_copy['LAST_LOGON_DATE_DATEONLY'] > df_copy['TERMINATION_DATE_DATEONLY'])]
            results = []
            if len(filtered_df_1) > 0:
                results.append({
                    'status': 'warning',
                    'title': 'Terminated but Enabled',
                    'message': f'{len(filtered_df_1)} istifadəçi ixtisar edilib amma hələ aktivdir'
                })
            if len(filtered_df_2) > 0:
                results.append({
                    'status': 'error',
                    'title': 'Logon After Termination',
                    'message': f'{len(filtered_df_2)} istifadəçi ixtisardan sonra daxil olub'
                })
            return {'success': True, 'details': results, 'data': {
                'terminated_enabled': filtered_df_1.to_dict('records'),
                'logon_after_term': filtered_df_2.to_dict('records')
            }}
        except Exception as e:
            return {'success': False, 'error': str(e)}
    
    def vendor_access_monitoring(self):
        try:
            filtered_df = self.df[
                (self.df['ENABLED'] == True) &
                (self.df['PO_BOX'] == 'V') &
                (self.df['DESCRIPTION'].astype(str).str.contains("SD-", na=False)) &
                (self.df['ACCOUNT_EXPIRATION_DATE'].isna())
            ]
            results = [{
                'status': 'warning' if len(filtered_df) > 0 else 'success',
                'title': 'Vendor Access Check',
                'message': f'{len(filtered_df)} vendor hesabının müddəti təyin edilməyib'
            }]
            return {'success': True, 'details': results, 'data': filtered_df.to_dict('records')}
        except Exception as e:
            return {'success': False, 'error': str(e)}
    
    def procedure_alignment_check(self):
        try:
            results = []
            excluded_values = ['S', 'X', 'MR', 'RM']
            df_main = self.df[~self.df['PO_BOX'].astype(str).str.strip().str.upper().isin(excluded_values)]
            not_required_df = df_main[
                df_main['PASSWORD_NOT_REQUIRED'].notna() &
                (df_main['PASSWORD_NOT_REQUIRED'].astype(str).str.strip() != '')
            ]
            df_main['PASSWORD_LAST_SET'] = pd.to_datetime(df_main['PASSWORD_LAST_SET'], errors='coerce')
            today = pd.to_datetime(datetime.today().date())
            df_main['DAY_DIFF'] = (today - df_main['PASSWORD_LAST_SET']).dt.days
            password_check_df = df_main[
                (df_main['DAY_DIFF'] >= 60) &
                (df_main['ENABLED'] == True) &
                (df_main['PASSWORD_EXPIRED'] == False) &
                (df_main['PASSWORD_NEVER_EXPIRES'] == False)
            ]
            if len(not_required_df) > 0:
                results.append({
                    'status': 'error',
                    'title': 'Password Not Required',
                    'message': f'{len(not_required_df)} hesabda şifrə tələb edilmir'
                })
            if len(password_check_df) > 0:
                results.append({
                    'status': 'warning',
                    'title': 'Old Passwords',
                    'message': f'{len(password_check_df)} hesabın şifrəsi 60 gündən köhnədir'
                })
            return {'success': True, 'details': results, 'data': {
                'password_not_required': not_required_df.to_dict('records'),
                'old_passwords': password_check_df.to_dict('records')
            }}
        except Exception as e:
            return {'success': False, 'error': str(e)}

class TWOChecker:
    def __init__(self, excel_file_path):
        self.excel_file = excel_file_path
        self.df = pd.read_excel(excel_file_path)
    
    def access_termination(self):
        try:
            results = []
            condition_nonopen = (
                self.df['TERMINATION_DATE'].notna() & 
                self.df['GROUPNAME'].notna() & 
                (self.df['GROUPNAME'].astype(str).str.strip() != '') & 
                (~self.df['GROUPNAME'].astype(str).str.lower().str.contains('blocked', na=False)) &
                (self.df['STATUS'].astype(str).str.upper() != 'OPEN') &
                (self.df['HR_CODE'].astype(str).str.upper().str.strip() != 'S')
            )
            filtered_nonopen_df = self.df[condition_nonopen]
            df_copy = self.df.copy()
            df_copy['TERMINATION_DATE'] = pd.to_datetime(df_copy['TERMINATION_DATE'], errors='coerce')
            filtered_active_df = df_copy[
                df_copy['TERMINATION_DATE'].notna() &
                (df_copy['STATUS'].str.upper() == 'ACTIVE')
            ]
            results.append({
                'status': 'success',
                'title': 'Non-Open Users',
                'message': f'{len(filtered_nonopen_df)} istifadəçi tapıldı'
            })
            if len(filtered_active_df) > 0:
                results.append({
                    'status': 'warning',
                    'title': 'Active Users (İxtisar edilmiş)',
                    'message': f'{len(filtered_active_df)} istifadəçi hələ ACTIVE statusundadır'
                })
            return {'success': True, 'details': results, 'data': {
                'non_open': filtered_nonopen_df.to_dict('records'),
                'active_users': filtered_active_df.to_dict('records')
            }}
        except Exception as e:
            return {'success': False, 'error': str(e)}
    
    def user_workplace_match(self):
        try:
            filtered_df = self.df[
                self.df['NAME'].astype(str).str.startswith(('BO', 'HO'), na=False) &
                (self.df['STATUS'].astype(str).str.strip() == 'Active') &
                (self.df['CURRENT_POSITION'].isna() |
                 (self.df['CURRENT_POSITION'].astype(str).str.strip() == '') |
                 (self.df['CURRENT_POSITION'].astype(str).str.strip() == '/  /  /')) &
                self.df['HR_CODE'].astype(str).str.startswith('T', na=False)
            ]
            filtered_df = filtered_df.drop_duplicates(subset=['HR_CODE'], keep='first')
            results = [{
                'status': 'warning' if len(filtered_df) > 0 else 'success',
                'title': 'User-Workplace Match',
                'message': f'{len(filtered_df)} uyğunsuzluq tapıldı'
            }]
            return {'success': True, 'details': results, 'data': filtered_df.to_dict('records')}
        except Exception as e:
            return {'success': False, 'error': str(e)}


def get_checker(system_name, excel_file_path):
    system_lower = system_name.lower()
    if 'cms' in system_lower:
        return CMSChecker(excel_file_path)
    elif 'fimi' in system_lower:
        return FIMIChecker(excel_file_path)
    elif 'kapitalcard' in system_lower:
        return KapitalCardChecker(excel_file_path)
    elif 'kapitalho' in system_lower:
        return KapitalhoChecker(excel_file_path)
    elif 'two' in system_lower:
        return TWOChecker(excel_file_path)
    else:
        if system_name in ['Elma Docflow', 'Elma BPM', 'Odin', 'CBAPP', 'Optimus', 'Zeus']:
            return KapitalhoChecker(excel_file_path)
        elif system_name == 'SWIFT':
            return FIMIChecker(excel_file_path)
        else:
            raise ValueError(f"Unknown system: {system_name}")

@app.route('/api/execute-check', methods=['POST'])
def execute_check():
    if 'file' not in request.files:
        return jsonify({'success': False, 'error': 'Fayl tapılmadı'}), 400
    file = request.files['file']
    if file.filename == '':
        return jsonify({'success': False, 'error': 'Fayl seçilməyib'}), 400
    if not allowed_file(file.filename):
        return jsonify({'success': False, 'error': 'Yalnız Excel faylları (.xlsx, .xls) qəbul edilir'}), 400
    check_type = request.form.get('check_type')
    system_name = request.form.get('system_name')
    system_type = request.form.get('system_type')
    if not all([check_type, system_name, system_type]):
        return jsonify({'success': False, 'error': 'Bütün parametrlər daxil edilməlidir'}), 400
    try:
        filename = secure_filename(file.filename)
        filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
        file.save(filepath)
        checker = get_checker(system_name, filepath)
        result = None
        if check_type == 'access_termination':
            result = checker.access_termination()
        elif check_type == 'vacation_monitoring':
            if hasattr(checker, 'vacation_monitoring'):
                result = checker.vacation_monitoring()
            else:
                result = {'success': False, 'error': f'Vacation monitoring {system_name} üçün mövcud deyil'}
        elif check_type == 'user_workplace_match':
            if hasattr(checker, 'user_workplace_match'):
                result = checker.user_workplace_match()
            else:
                result = {'success': False, 'error': f'User-workplace match {system_name} üçün mövcud deyil'}
        elif check_type == 'expiry_date_monitoring':
            if hasattr(checker, 'expiry_date_monitoring'):
                result = checker.expiry_date_monitoring()
            else:
                result = {'success': False, 'error': f'Expiry date monitoring {system_name} üçün mövcud deyil'}
        elif check_type == 'vendor_access_monitoring':
            if hasattr(checker, 'vendor_access_monitoring'):
                result = checker.vendor_access_monitoring()
            else:
                result = {'success': False, 'error': f'Vendor monitoring {system_name} üçün mövcud deyil'}
        elif check_type == 'procedure_alignment':
            if hasattr(checker, 'procedure_alignment_check'):
                result = checker.procedure_alignment_check()
            else:
                result = {'success': False, 'error': f'Procedure alignment {system_name} üçün mövcud deyil'}
        elif check_type == 'rotation_monitoring':
            result = {'success': False, 'error': 'Rotation monitoring hazırda implementasiya edilməyib'}
        elif check_type == 'new_user_access':
            result = {'success': False, 'error': 'New user access provision hazırda implementasiya edilməyib'}
        else:
            result = {'success': False, 'error': 'Naməlum yoxlama növü'}
        os.remove(filepath)
        if result['success']:
            result['metadata'] = {
                'check_type': check_type,
                'system_name': system_name,
                'system_type': system_type,
                'timestamp': datetime.now().isoformat()
            }
        return jsonify(result)
    except Exception as e:
        if 'filepath' in locals() and os.path.exists(filepath):
            os.remove(filepath)
        return jsonify({'success': False, 'error': f'Xəta baş verdi: {str(e)}'}), 500

@app.route('/api/export-results', methods=['POST'])
def export_results():
    try:
        data = request.json
        results_data = data.get('data', {})
        check_type = data.get('check_type', 'results')
        system_name = data.get('system_name', 'system')
        output_filename = f"{system_name}_{check_type}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        output_path = os.path.join(app.config['UPLOAD_FOLDER'], output_filename)
        with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
            for sheet_name, sheet_data in results_data.items():
                if sheet_data:
                    df = pd.DataFrame(sheet_data)
                    df.to_excel(writer, sheet_name=sheet_name[:31], index=False)
                    worksheet = writer.sheets[sheet_name[:31]]
                    adjust_column_widths(worksheet)
        with open(output_path, 'rb') as f:
            import base64
            file_content = base64.b64encode(f.read()).decode('utf-8')
        os.remove(output_path)
        return jsonify({'success': True, 'filename': output_filename, 'content': file_content})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/health', methods=['GET'])
def health_check():
    return jsonify({'status': 'healthy','timestamp': datetime.now().isoformat(),'version': '2.0'})

@app.route('/api/systems', methods=['GET'])
def get_systems():
    systems = [
        {'name': 'Elma Docflow', 'type': 'AD', 'integrated': True},
        {'name': 'Elma BPM', 'type': 'AD', 'integrated': True},
        {'name': 'Odin', 'type': 'AD', 'integrated': True},
        {'name': 'CBAPP', 'type': 'AD', 'integrated': True},
        {'name': 'Optimus', 'type': 'AD', 'integrated': True},
        {'name': 'Zeus', 'type': 'AD', 'integrated': True},
        {'name': 'Kapitalho', 'type': 'AD', 'integrated': True},
        {'name': 'Kapitalcard', 'type': 'AD', 'integrated': True},
        {'name': 'CMS', 'type': 'nAD', 'integrated': False},
        {'name': 'TWO', 'type': 'nAD', 'integrated': False},
        {'name': 'FIMI', 'type': 'nAD', 'integrated': False},
        {'name': 'SWIFT', 'type': 'nAD', 'integrated': False}
    ]
    return jsonify(systems)

@app.route('/api/checks', methods=['GET'])
def get_checks():
    checks = [
        {'id': 'access_termination','name': 'Access Termination','systems': ['CMS', 'FIMI', 'Kapitalcard', 'Kapitalho', 'TWO'],'applicableToAD': True,'applicableToNAD': True},
        {'id': 'vacation_monitoring','name': 'Vacation Monitoring','systems': ['CMS', 'FIMI'],'applicableToAD': False,'applicableToNAD': True},
        {'id': 'user_workplace_match','name': 'User-Work Place Match','systems': ['CMS', 'FIMI', 'TWO'],'applicableToAD': False,'applicableToNAD': True},
        {'id': 'expiry_date_monitoring','name': 'Expiry Date Monitoring','systems': ['CMS'],'applicableToAD': False,'applicableToNAD': True},
        {'id': 'vendor_access_monitoring','name': 'Vendor Access Monitoring','systems': ['Kapitalcard', 'Kapitalho'],'applicableToAD': True,'applicableToNAD': False},
        {'id': 'procedure_alignment','name': 'Procedure Alignment Check','systems': ['Kapitalcard', 'Kapitalho'],'applicableToAD': True,'applicableToNAD': False}
    ]
    return jsonify(checks)

if __name__ == '__main__':
    app.run(debug=True, port=5000, host='0.0.0.0')